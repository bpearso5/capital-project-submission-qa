"""Builds the synthetic data and the workbook. Seeded, so output is identical every run.
Usage: python3 scripts/build_workbook.py"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

random.seed(42)
F = "Arial"
BLACK = PatternFill("solid", fgColor="000000")
GRAY = PatternFill("solid", fgColor="E7E6E6")
INPUT = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
H = Font(name=F, bold=True, color="FFFFFF", size=10)
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True)
SUB = Font(name=F, size=10, italic=True, color="595959")

def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = H, BLACK, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def body(ws, r1, r2, ncol):
    for r in range(r1, r2 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            if cell.font != BB:
                cell.font = B

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

def title(ws, t, s):
    ws["A1"] = t; ws["A1"].font = TITLE
    ws["A2"] = s; ws["A2"].font = SUB

wb = Workbook()

# ---------------- README ----------------
ws = wb.active; ws.title = "README"
title(ws, "Capital Project Submission QA & Weekly Executive Report",
      "Recreated work sample - Brystal Pearson, CAPM - ALL DATA IS SYNTHETIC")
lines = [
    ("What this is", "A rebuilt version of work I did as a contract Project Analyst on a multi-region healthcare capital portfolio: routing invoices, POs, and commitment changes for approval, finding why submissions kept getting rejected, fixing it with a standardized pre-submission checklist, and reporting budget position to leadership weekly."),
    ("Data", "Every project, person, amount, and date is invented. Rejection categories are modeled on the kinds of issues I saw; no client records, files, or screenshots are included."),
    ("The story in one line", "I stopped fixing rejected submissions one at a time and fixed the template they were built from."),
    ("", ""),
    ("Tab", "Purpose"),
    ("Submission_Log", "96 synthetic submissions over 12 weeks. Weeks 1-6 = before the checklist; weeks 7-12 = after."),
    ("Rejection_Pareto", "Rejections by reason, before vs. after, with cumulative % - the analysis that pointed to the fix."),
    ("PreSubmit_Checklist", "The standardized checklist. Each item is tied to the rejection reason it prevents."),
    ("Budget_Status", "12 capital projects: approved budget, committed, invoiced, remaining, closeout date, and an automatic flag."),
    ("Weekly_Exec_Report", "One-page leadership summary. Change the report date (gray cell) and every figure recalculates."),
    ("", ""),
    ("Legend", "Gray-filled cells are inputs you can edit. Everything else is a formula or synthetic source data."),
    ("Note on results", "Before/after figures here come from synthetic data and illustrate the method. They are not measured outcomes from the engagement."),
]
for i, (a, b) in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=a).font = BB
    ws.cell(row=i, column=2, value=b).font = B
    ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
for c in (1, 2):
    ws.cell(row=8, column=c).font, ws.cell(row=8, column=c).fill = H, BLACK
widths(ws, [24, 100])

# ---------------- Projects ----------------
regions = ["Central", "East", "West"]
ptypes = ["OR Renovation", "Imaging Suite", "HVAC Replacement", "ED Expansion",
          "Roof Replacement", "Pharmacy Upgrade", "Generator Install", "Lab Remodel",
          "Parking Structure", "Patient Tower Refresh", "Sterile Processing", "Clinic Build-Out"]
owners = ["J. Alvarez", "K. Brooks", "M. Chen", "D. Okafor", "S. Patel"]
report_date = dt.date(2025, 11, 21)
projects = []
for i, t in enumerate(ptypes):
    pid = f"CP-{2401 + i}"
    approved = random.choice([350, 480, 620, 750, 900, 1200, 1500, 2100, 2800]) * 1000
    committed = round(approved * random.uniform(0.78, 1.0), -3)
    invoiced = round(committed * random.uniform(0.55, 1.0), -3)
    status = random.choice(["Active", "Active", "Active", "Closeout", "Closeout", "Closed"])
    target = report_date + dt.timedelta(days=random.randint(25, 200))
    if status == "Closed":
        invoiced = committed; target = report_date - dt.timedelta(days=random.randint(10, 60))
    projects.append([pid, f"{regions[i % 3]} - {t}", regions[i % 3], random.choice(owners),
                     approved, committed, invoiced, status, target])
# force a couple of interesting cases
projects[3][5] = projects[3][4] * 0.99; projects[3][7] = "Active"
projects[6][7] = "Closeout"; projects[6][8] = report_date - dt.timedelta(days=41)
projects[9][7] = "Closeout"; projects[9][8] = report_date - dt.timedelta(days=18)

# ---------------- Submission log ----------------
reasons = [
    ("R1", "Cost code missing or not on approved budget line"),
    ("R2", "Amount exceeds remaining commitment balance"),
    ("R3", "PO number missing or does not match commitment"),
    ("R4", "Required backup missing (signed pay app, lien waiver)"),
    ("R5", "Routed to wrong approver"),
    ("R6", "Retainage or math error"),
]
pre_w = [34, 22, 16, 14, 8, 6]   # pre-checklist reason mix
doc_types = ["Invoice", "Invoice", "Invoice", "PO", "Commitment Change"]
subs = []
n = 0
start = dt.date(2025, 9, 1)
for wk in range(1, 13):
    for _ in range(8):
        n += 1
        p = random.choice(projects[:11])
        doc = random.choice(doc_types)
        amt = round(random.uniform(4, 180) * 1000, -2)
        pre = wk <= 6
        rej = random.random() < (0.36 if pre else 0.09)
        reason = ""
        if rej:
            if pre:
                reason = random.choices(reasons, weights=pre_w)[0][1]
            else:  # residual: mostly things a checklist can't fully stop
                reason = random.choices(reasons, weights=[6, 30, 8, 10, 30, 16])[0][1]
        days = random.randint(6, 14) if rej else random.randint(2, 5)
        subs.append([f"S-{1000 + n}", wk, start + dt.timedelta(weeks=wk - 1, days=random.randint(0, 4)),
                     p[0], doc, amt, random.choice(owners),
                     "No" if pre else "Yes", "Rejected" if rej else "Approved", reason,
                     1 if rej else 0, days])

ws = wb.create_sheet("Submission_Log")
title(ws, "Submission Log", "Synthetic. Checklist introduced at week 7.")
cols = ["Submission ID", "Week", "Submitted", "Project ID", "Doc Type", "Amount ($)", "Submitted By",
        "Checklist Used", "First-Pass Outcome", "Rejection Reason", "Resubmissions", "Days to Final Approval"]
header(ws, 4, cols)
for r, row in enumerate(subs, 5):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
    ws.cell(row=r, column=3).number_format = "mm/dd/yyyy"
    ws.cell(row=r, column=6).number_format = "$#,##0"
LAST = 4 + len(subs)
body(ws, 5, LAST, len(cols))
ws.auto_filter.ref = f"A4:L{LAST}"
ws.freeze_panes = "A5"
widths(ws, [13, 7, 11, 11, 18, 12, 13, 10, 12, 50, 9, 11])
ws.row_dimensions[4].height = 30
SL = "Submission_Log"
rng = lambda col: f"{SL}!${col}$5:${col}${LAST}"

# ---------------- Rejection Pareto ----------------
ws = wb.create_sheet("Rejection_Pareto")
title(ws, "Rejection Pareto - Before vs. After Checklist", "Before = weeks 1-6 (no checklist). After = weeks 7-12.")
header(ws, 4, ["Code", "Rejection Reason", "Before", "% of Before", "Cumulative % (Before)", "After", "Change"])
from collections import Counter
cnt = Counter(x[9] for x in subs if x[7] == "No" and x[8] == "Rejected")
order = sorted(range(6), key=lambda i: (-cnt[reasons[i][1]], i))  # true Pareto order
for k, i in enumerate(order):
    r = 5 + k
    ws.cell(row=r, column=1, value=reasons[i][0])
    ws.cell(row=r, column=2, value=reasons[i][1])
    ws.cell(row=r, column=3, value=f'=COUNTIFS({rng("J")},B{r},{rng("H")},"No")')
    ws.cell(row=r, column=4, value=f"=IFERROR(C{r}/$C$11,0)")
    ws.cell(row=r, column=5, value=f"=SUM($D$5:D{r})")
    ws.cell(row=r, column=6, value=f'=COUNTIFS({rng("J")},B{r},{rng("H")},"Yes")')
    ws.cell(row=r, column=7, value=f"=F{r}-C{r}")
ws["B11"] = "Total rejections"; ws["B11"].font = BB
for col in "CF":
    ws[f"{col}11"] = f"=SUM({col}5:{col}10)"; ws[f"{col}11"].font = BB
ws["G11"] = "=F11-C11"; ws["G11"].font = BB
body(ws, 5, 11, 7)
for r in range(5, 12):
    for col in "DE":
        ws[f"{col}{r}"].number_format = "0.0%"
    ws[f"G{r}"].number_format = "+0;-0;0"
ws["E11"] = None
ws["A13"] = "Read it this way:"; ws["A13"].font = BB
ws["A14"] = "The top three reasons by count were preventable before submission - they were template problems, not people problems. That is what the checklist targets."
ws["A15"] = "R2 and R5 persist after: budget balance and approver changes happen upstream, so they need a budget refresh and a routing matrix, not just a checklist."
for a in ("A14", "A15"): ws[a].font = B
widths(ws, [8, 52, 10, 12, 14, 10, 10])
ws.row_dimensions[4].height = 30

# ---------------- Checklist ----------------
ws = wb.create_sheet("PreSubmit_Checklist")
title(ws, "Pre-Submission Checklist", "Completed by the submitter before routing. Analyst confirms before it moves forward.")
header(ws, 4, ["#", "Check", "How to verify", "Prevents", "Done (Y/N)"])
items = [
    ("Cost code matches an approved budget line", "Compare to current budget line list for the project", "R1"),
    ("Amount is within the remaining commitment balance", "Remaining = commitment - invoiced to date; pull from Budget_Status", "R2"),
    ("PO number present and matches the commitment", "Match PO on document to commitment record", "R3"),
    ("Signed pay application attached", "Contractor and architect signatures present", "R4"),
    ("Conditional lien waiver attached for this period", "Period and amount match the pay application", "R4"),
    ("Approver matches the routing matrix for this amount and region", "Check threshold tier before routing", "R5"),
    ("Retainage calculated at contract rate", "Recalculate; line totals tie to the summary", "R6"),
    ("Document type and project ID are correct", "Matches project record exactly", "R1"),
]
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
ws.add_data_validation(dv)
for k, (a, b, c) in enumerate(items):
    r = 5 + k
    for col, v in enumerate([k + 1, a, b, c, ""], 1):
        ws.cell(row=r, column=col, value=v)
    ws.cell(row=r, column=5).fill = INPUT
    dv.add(f"E{r}")
body(ws, 5, 12, 5)
ws["B14"] = "Items complete"; ws["B14"].font = BB
ws["E14"] = '=COUNTIF(E5:E12,"Y")&" of "&COUNTA(B5:B12)'
ws["B15"] = "Ready to route?"; ws["B15"].font = BB
ws["E15"] = '=IF(COUNTIF(E5:E12,"Y")=COUNTA(B5:B12),"YES - route","NO - hold")'
for a in ("E14", "E15"): ws[a].font = BB
ws["B17"] = "Example row: enter Y or N in the gray cells. All eight must be Y before a submission routes."
ws["B17"].font = SUB
ws["E5"] = "Y"
for col in range(2, 4):
    for r in range(5, 13):
        ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
widths(ws, [5, 48, 52, 10, 11])

# ---------------- Budget status ----------------
ws = wb.create_sheet("Budget_Status")
title(ws, "Capital Project Budget Status", "Synthetic. Amounts in dollars.")
ws["A3"] = "Report date:"; ws["A3"].font = BB
ws["B3"] = "=Weekly_Exec_Report!C4"; ws["B3"].number_format = "mm/dd/yyyy"; ws["B3"].font = Font(name=F, size=10, color="008000")
cols = ["Project ID", "Project", "Region", "Owner", "Approved Budget ($)", "Committed ($)", "Invoiced to Date ($)",
        "Uncommitted Budget ($)", "Open Commitment ($)", "% Committed", "Status", "Target Closeout", "Flag"]
header(ws, 5, cols)
for k, p in enumerate(projects):
    r = 6 + k
    vals = p[:7]
    for c, v in enumerate(vals, 1):
        ws.cell(row=r, column=c, value=v)
    ws.cell(row=r, column=8, value=f"=E{r}-F{r}")
    ws.cell(row=r, column=9, value=f"=F{r}-G{r}")
    ws.cell(row=r, column=10, value=f"=IFERROR(F{r}/E{r},0)")
    ws.cell(row=r, column=11, value=p[7])
    ws.cell(row=r, column=12, value=p[8])
    ws.cell(row=r, column=13, value=(f'=IF(K{r}="Closed","",IF(AND(L{r}<$B$3,K{r}<>"Closed"),"Closeout overdue",'
                                      f'IF(J{r}>=0.95,"Over 95% committed",IF(L{r}-$B$3<=30,"Closeout within 30 days",""))))'))
    for c in range(5, 10):
        ws.cell(row=r, column=c).number_format = "$#,##0;($#,##0);-"
    ws.cell(row=r, column=10).number_format = "0.0%"
    ws.cell(row=r, column=12).number_format = "mm/dd/yyyy"
PL = 5 + len(projects)
ws.cell(row=PL + 1, column=2, value="Portfolio total").font = BB
for c in range(5, 10):
    L = get_column_letter(c)
    cell = ws.cell(row=PL + 1, column=c, value=f"=SUM({L}6:{L}{PL})")
    cell.font = BB; cell.number_format = "$#,##0;($#,##0);-"
ws.cell(row=PL + 1, column=10, value=f"=IFERROR(F{PL+1}/E{PL+1},0)").number_format = "0.0%"
body(ws, 6, PL + 1, len(cols))
ws.conditional_formatting.add(f"A6:M{PL}", FormulaRule(formula=[f'$M6<>""'], font=Font(name=F, bold=True)))
ws.conditional_formatting.add(f"M6:M{PL}", FormulaRule(formula=[f'$M6="Closeout overdue"'], fill=GRAY, font=Font(name=F, bold=True)))
ws.freeze_panes = "C6"
widths(ws, [10, 30, 9, 11, 14, 13, 13, 13, 13, 10, 10, 11, 22])
ws.row_dimensions[5].height = 30

# ---------------- Weekly exec report ----------------
ws = wb.create_sheet("Weekly_Exec_Report")
title(ws, "Weekly Executive Report - Capital Portfolio", "Synthetic data. One page, decision-ready.")
ws["A4"] = "Report date"; ws["A4"].font = BB
ws["C4"] = report_date; ws["C4"].fill = INPUT; ws["C4"].number_format = "mm/dd/yyyy"; ws["C4"].font = BB
ws["A5"] = "Report week"; ws["A5"].font = BB
ws["C5"] = "=MIN(12,MAX(1,INT((C4-DATE(2025,9,1))/7)+1))"; ws["C5"].font = BB
ws["D5"] = "Week 1 starts 09/01/2025; capped at week 12 (end of sample data)."; ws["D5"].font = SUB

ws["A7"] = "1. BUDGET POSITION"; ws["A7"].font = BB
bs = "Budget_Status"
kp = [
    ("Approved budget", f"={bs}!E{PL+1}", "$#,##0"),
    ("Committed", f"={bs}!F{PL+1}", "$#,##0"),
    ("Invoiced to date", f"={bs}!G{PL+1}", "$#,##0"),
    ("Uncommitted budget", f"={bs}!H{PL+1}", "$#,##0"),
    ("% of budget committed", f"={bs}!J{PL+1}", "0.0%"),
]
header(ws, 8, ["Measure", "", "Value"])
for k, (a, f, nf) in enumerate(kp):
    r = 9 + k
    ws.cell(row=r, column=1, value=a); ws.cell(row=r, column=3, value=f).number_format = nf
    ws.cell(row=r, column=3).font = Font(name=F, size=10, color="008000")
body(ws, 9, 13, 3)

ws["A15"] = "2. SUBMISSION QUALITY"; ws["A15"].font = BB
header(ws, 16, ["Measure", "Before checklist (wk 1-6)", "After checklist (wk 7-12)", "This week"])
wkr = "$C$5"
q = [
    ("Submissions",
     f'=COUNTIFS({rng("H")},"No")', f'=COUNTIFS({rng("H")},"Yes")', f'=COUNTIFS({rng("B")},{wkr})', "0"),
    ("Rejected on first pass",
     f'=COUNTIFS({rng("H")},"No",{rng("I")},"Rejected")', f'=COUNTIFS({rng("H")},"Yes",{rng("I")},"Rejected")',
     f'=COUNTIFS({rng("B")},{wkr},{rng("I")},"Rejected")', "0"),
    ("First-pass approval rate", "=IFERROR(1-B18/B17,0)", "=IFERROR(1-C18/C17,0)", "=IFERROR(1-D18/D17,0)", "0.0%"),
    ("Avg days to final approval",
     f'=AVERAGEIFS({rng("L")},{rng("H")},"No")', f'=AVERAGEIFS({rng("L")},{rng("H")},"Yes")',
     f'=IFERROR(AVERAGEIFS({rng("L")},{rng("B")},{wkr}),0)', "0.0"),
]
for k, (a, b, c, d, nf) in enumerate(q):
    r = 17 + k
    ws.cell(row=r, column=1, value=a)
    for col, v in zip((2, 3, 4), (b, c, d)):
        ws.cell(row=r, column=col, value=v).number_format = nf
body(ws, 17, 20, 4)
ws["A21"] = "Top rejection reason before checklist:"; ws["A21"].font = BB
ws["C21"] = "=INDEX(Rejection_Pareto!B5:B10,MATCH(MAX(Rejection_Pareto!C5:C10),Rejection_Pareto!C5:C10,0))"
ws["A22"] = "Top rejection reason after checklist:"; ws["A22"].font = BB
ws["C22"] = "=INDEX(Rejection_Pareto!B5:B10,MATCH(MAX(Rejection_Pareto!F5:F10),Rejection_Pareto!F5:F10,0))"

ws["A24"] = "3. PROJECTS NEEDING ATTENTION"; ws["A24"].font = BB
ws["A25"] = "Flagged projects:"; ws["A25"].font = BB
ws["C25"] = f'=COUNTIF({bs}!M6:M{PL},"?*")'; ws["C25"].font = BB
header(ws, 26, ["Project ID", "Project", "Owner", "Flag", "Open Commitment ($)"])
for k in range(len(projects)):
    r = 27 + k; s = 6 + k
    ws.cell(row=r, column=1, value=f"={bs}!A{s}")
    ws.cell(row=r, column=2, value=f"={bs}!B{s}")
    ws.cell(row=r, column=3, value=f"={bs}!D{s}")
    ws.cell(row=r, column=4, value=f'=IF({bs}!M{s}="","-",{bs}!M{s})')
    ws.cell(row=r, column=5, value=f"={bs}!I{s}").number_format = "$#,##0;($#,##0);-"
RE = 26 + len(projects)
body(ws, 27, RE, 5)
ws.conditional_formatting.add(f"A27:E{RE}", FormulaRule(formula=['$D27<>"-"'], font=Font(name=F, bold=True), fill=GRAY))

ws[f"A{RE+2}"] = "4. DECISIONS / ASKS FOR LEADERSHIP"; ws[f"A{RE+2}"].font = BB
asks = ["Approve a quarterly budget-line refresh so R2 (over remaining balance) stops recurring.",
        "Publish an approver routing matrix by amount and region to close out R5.",
        "Confirm owners and revised dates for every project flagged 'Closeout overdue'."]
for k, a in enumerate(asks):
    ws.cell(row=RE + 3 + k, column=1, value=f"{k+1}. {a}").font = B
widths(ws, [30, 34, 26, 24, 18])
ws.row_dimensions[16].height = 30
ws.print_area = f"A1:E{RE+5}"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.move_sheet("Weekly_Exec_Report", offset=-4)
import csv, os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
with open(os.path.join(ROOT, "data", "submission_log.csv"), "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["submission_id","week","submitted","project_id","doc_type","amount","submitted_by",
                "checklist_used","first_pass_outcome","rejection_reason","resubmissions","days_to_final_approval"])
    for x in subs: w.writerow(x)
with open(os.path.join(ROOT, "data", "projects.csv"), "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["project_id","project","region","owner","approved_budget","committed","invoiced_to_date","status","target_closeout"])
    for p in projects: w.writerow(p)
out = os.path.join(ROOT, "workbook", "Capital_Project_Submission_QA_and_Weekly_Report.xlsx")
wb.save(out)
print(out, LAST, PL)
